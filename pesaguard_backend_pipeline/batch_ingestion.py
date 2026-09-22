"""Durable file-import jobs built on top of the provider ingestion boundary."""

from __future__ import annotations

import csv
import io
import json
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, Mapping, Optional

from ingestion import IngestionError, IngestionService
from models import ImportJob
from partitioning import raw_object_partition


SUPPORTED_FORMATS = {"csv", "json", "xlsx", "parquet"}
MAX_IMPORT_BYTES = int(os.getenv("PESAGUARD_IMPORT_MAX_BYTES", str(50 * 1024 * 1024)))
MAX_IMPORT_RECORDS = int(os.getenv("PESAGUARD_IMPORT_MAX_RECORDS", "250000"))


class BatchImportError(ValueError):
    """Raised when an import cannot be safely submitted or parsed."""


class ObjectStorage:
    """Small local object-store boundary used by development and tests.

    Production deployments can replace this class with an S3-compatible
    implementation without changing import job or parsing behavior.
    """

    def __init__(self, root: Optional[str] = None):
        self.root = Path(root or os.getenv("PESAGUARD_IMPORT_OBJECT_STORAGE", "var/imports"))

    def put(self, tenant_id: str, job_id: str, filename: str, content: bytes, *, source: str = "unknown", observed_at: Any = None) -> str:
        object_key = f"{raw_object_partition(tenant_id, source, observed_at)}/{job_id}/{filename}"
        destination = self.root / object_key
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        return object_key

    def get(self, object_key: str) -> bytes:
        path = self.root / object_key
        if not path.is_file():
            raise BatchImportError("import object is not available")
        return path.read_bytes()


def _format_for(filename: str) -> str:
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_FORMATS:
        raise BatchImportError(f"unsupported import format: {suffix or 'unknown'}")
    return suffix


def _safe_filename(filename: str) -> str:
    name = Path(filename or "").name
    if not name or name in {".", ".."} or len(name) > 255:
        raise BatchImportError("a valid filename is required")
    return name


def _records_from_bytes(filename: str, content: bytes) -> Iterator[Mapping[str, Any]]:
    file_format = _format_for(filename)
    if file_format == "csv":
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        if not reader.fieldnames:
            raise BatchImportError("CSV file must contain a header row")
        yield from reader
        return
    if file_format == "json":
        try:
            document = json.loads(content.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise BatchImportError("JSON file is invalid") from exc
        records = document.get("records") if isinstance(document, dict) else document
        if not isinstance(records, list):
            raise BatchImportError("JSON file must contain an array or a records array")
        for record in records:
            if not isinstance(record, Mapping):
                raise BatchImportError("JSON records must be objects")
            yield record
        return
    if file_format == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise BatchImportError("XLSX imports require the openpyxl package") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
            if not headers or not any(headers):
                raise BatchImportError("XLSX file must contain a header row")
            for values in rows:
                yield {header: value for header, value in zip(headers, values) if header}
        finally:
            workbook.close()
        return
    try:
        import pyarrow.parquet as parquet
    except ImportError as exc:
        raise BatchImportError("Parquet imports require the pyarrow package") from exc
    table = parquet.read_table(io.BytesIO(content))
    yield from table.to_pylist()


class BatchImportService:
    """Submit and process bounded files through the shared ingestion service."""

    def __init__(self, ingestion_service: IngestionService, storage: Optional[ObjectStorage] = None):
        self.ingestion_service = ingestion_service
        self.storage = storage or ObjectStorage()

    def submit(
        self,
        session: Any,
        *,
        tenant_id: str,
        source: str,
        filename: str,
        content: bytes,
    ) -> ImportJob:
        tenant = str(tenant_id or "").strip()
        if not tenant:
            raise BatchImportError("tenant context is required")
        if not content or len(content) > MAX_IMPORT_BYTES:
            raise BatchImportError("import file is empty or exceeds the size limit")
        safe_name = _safe_filename(filename)
        file_format = _format_for(safe_name)
        provider = str(source or "").strip().lower()
        if provider not in self.ingestion_service.adapters:
            raise BatchImportError(f"unsupported import source: {source}")
        job_id = f"imp_{uuid.uuid4().hex}"
        object_key = self.storage.put(tenant, job_id, safe_name, content, source=provider)
        job = ImportJob(
            id=job_id,
            tenant_id=tenant,
            source=provider,
            filename=safe_name,
            file_format=file_format,
            object_key=object_key,
            status="queued",
        )
        session.add(job)
        session.commit()
        return job

    def process(self, session: Any, import_id: str, *, tenant_id: Optional[str] = None) -> ImportJob:
        query = session.query(ImportJob).filter(ImportJob.id == import_id)
        if tenant_id is not None:
            query = query.filter(ImportJob.tenant_id == tenant_id)
        job = query.first()
        if job is None:
            raise BatchImportError("import job not found")
        if job.status == "completed":
            return job
        job.status = "running"
        job.started_at = datetime.now(timezone.utc)
        session.commit()
        errors = list(job.error_summary or [])
        try:
            records = _records_from_bytes(job.filename, self.storage.get(job.object_key))
            for record in records:
                if job.records_received >= MAX_IMPORT_RECORDS:
                    raise BatchImportError("import exceeds the maximum record limit")
                job.records_received += 1
                try:
                    result = self.ingestion_service.ingest(job.source, record, tenant_id=job.tenant_id)
                    if result.result.value in {"stored", "duplicate"}:
                        job.records_valid += 1
                    else:
                        job.records_failed += 1
                        if len(errors) < 20:
                            errors.append({"record": job.records_received, "error": "transaction persistence failed"})
                except (IngestionError, ValueError, TypeError) as exc:
                    job.records_failed += 1
                    if len(errors) < 20:
                        errors.append({"record": job.records_received, "error": str(exc)[:500]})
                if job.records_received % 100 == 0:
                    session.commit()
            job.status = "completed" if job.records_failed == 0 else "failed"
        except Exception as exc:
            job.status = "failed"
            if len(errors) < 20:
                errors.append({"record": job.records_received, "error": str(exc)[:500]})
        job.error_summary = errors
        job.completed_at = datetime.now(timezone.utc)
        session.commit()
        return job


def process_batch_import_job(import_id: str, database_url: Optional[str] = None) -> Dict[str, Any]:
    """Worker entry point for RQ/cron scheduled imports."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from event_store import EventStore

    url = database_url or os.getenv("DATABASE_URL")
    if not url:
        raise BatchImportError("DATABASE_URL is required for batch processing")
    engine = create_engine(url, connect_args={"check_same_thread": False} if url.startswith("sqlite") else {})
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    with session_factory() as session:
        service = BatchImportService(IngestionService(EventStore(database_url=url)))
        job = service.process(session, import_id)
        return {"import_id": job.id, "status": job.status, "records_received": job.records_received}


def process_queued_imports(database_url: Optional[str] = None, limit: int = 10) -> Dict[str, Any]:
    """Bounded scheduler entry point for daily and historical import runs."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from event_store import EventStore

    url = database_url or os.getenv("DATABASE_URL")
    if not url:
        raise BatchImportError("DATABASE_URL is required for scheduled imports")
    engine = create_engine(url, connect_args={"check_same_thread": False} if url.startswith("sqlite") else {})
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    processed = []
    with session_factory() as session:
        service = BatchImportService(IngestionService(EventStore(database_url=url)))
        jobs = session.query(ImportJob).filter(ImportJob.status == "queued").order_by(ImportJob.created_at).limit(max(1, min(int(limit), 100))).all()
        for queued in jobs:
            job = service.process(session, queued.id)
            processed.append({"import_id": job.id, "status": job.status})
    return {"processed": processed, "count": len(processed)}